import pandas as pd
from PIL import Image
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.chart import PieChart, PieChart3D
from openpyxl.chart import Reference, Series
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
loc = 'C:/Users/allen/OneDrive/Desktop/Github/Learning/' \
      'Python/Automation/Excel/data/2'                          # Location of the directory
wb = Workbook()                                                 # Create workbook
ws = wb.active                                                  # Go to active sheet of wb

data = [
    ['Flavour', 'Number Sold'],
    ['Chocolate', 120],
    ['Vanilla', 30],
    ['Strawberry', 20],
    ['Pistachio', 90],
    ['Mint Chocolate Chip', 40],
    ['Cookies and Cream', 60]
]                                                               # Data to be written
for z in data: ws.append(z)                                     # Write data to the sheet

# Chart
chart = PieChart()                                              # Create chart
chart.title = 'Ice_Cream_Chart'                                 # Title of the chart
labels = Reference(ws, min_col=1, min_row=2, max_row=7)         # Labels of the chart
data = Reference(ws, min_col=2, min_row=1, max_row=7)           # Data of the chart
chart.add_data(data, titles_from_data=True)                     # Add data to the chart
chart.set_categories(labels)                                    # Set labels of the chart

# Table
tab = Table(displayName='Ice_Cream_Table', ref='A1:B7')         # Create table
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium10',  # Name of the style from excel
    showFirstColumn=False, showLastColumn=False,
    showRowStripes=False, showColumnStripes=True
)                                                               # Style of the table   

# Image
img = Image(f'{loc}/C_Image.png')                               # Create image
img.width = img.width * 0.71                                    # Width of the image
img.height = img.height * 0.71                                  # Height of the image


# Save
ws.add_chart(chart, 'D1')                                       # Add chart to the sheet
ws.add_table(tab)                                               # Add table to the sheet
ws.add_image(img, 'N1')                                         # Add image to the sheet
wb.save(f'{loc}/C_Write.xlsx')                                  # Save the workbook